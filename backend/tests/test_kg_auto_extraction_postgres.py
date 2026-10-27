"""PostgreSQL integration tests: automatic research after a successful Word import.

What is proven here (all through the **real** import paths, never by calling the
extraction service directly to stand in for an import):

* creating a group starts nothing (no model call, no run);
* importing one Word into a group with auto research enabled creates a persistent run,
  delivers it to the queue, processes blocks, writes mentions/relationships/evidence
  and the group graph reflects them;
* a batch import queues **every** successfully imported document (no silent drops);
* a duplicate upload/re-import neither re-queues nor duplicates evidence;
* an unreachable queue keeps the run in the ledger with a reason and can be
  re-delivered;
* a new document does not re-extract the already-finished blocks of the old one;
* budget exhaustion pauses the automatic run and resuming needs a new ceiling;
* extraction only ever produces candidates: the official view stays empty until a
  human reviews, and evidence isolation between groups still holds;
* enabling the switch (or reading pages) never starts a historical backfill;
* attaching an already-uploaded Word to a group is itself a research event.

The language model is mocked at ``kg_llm_client.call_extraction_llm``; the parser,
writer, ledger and scope queries are the real ones.
"""
import unittest
from io import BytesIO
from unittest import mock
from unittest.mock import AsyncMock

from docx import Document as DocxDocument
from fastapi.testclient import TestClient

from app.test_support.bootstrap import require_isolated_test_environment

require_isolated_test_environment(__name__)

from app.api.routers import kg as kg_router  # noqa: F401 - ensures the router is loaded
from app.config import settings
from app.database import get_db
from app.main import app
from app.models.block import Block
from app.models.document import Document
from app.models.kg.enums import ClaimStatus, IdentityStatus, ReviewStatusKG
from app.models.kg.extraction_run import (
    KGExtractionItem,
    KGExtractionRun,
    KGExtractionRunTrigger,
)
from app.services.kg.llm_client import kg_llm_client  # noqa: E402
from app.models.kg.mention import PersonMention
from app.models.kg.person import Person
from app.models.kg.relationship import Relationship, RelationshipEvidence
from app.models.page import Page
from app.models.research_group import ResearchGroup, ResearchGroupDocument
from app.schemas.research_group import GroupCreate
from app.services.kg import auto_extraction, extraction_run_service
from app.services.research_group_service import ResearchGroupService
from app.test_support.db_fixtures import (
    make_session_factory,
    make_test_engine,
    requires_postgres_test_db,
)

SYNTH = "[SYNTH-AUTO]"
GROUP_MAIN = SYNTH + " Auto research group"
GROUP_OTHER = SYNTH + " Other research group"
GROUP_HISTORY = SYNTH + " History group"

EXTRACTION_JSON = """
{
  "entities": [
    {"mention": "%(prefix)s Henry Browne", "normalized_name": "%(prefix)s Henry Browne",
     "role": "supercargo", "organization": "East India Company",
     "mention_type": "person_name", "confidence": 0.9,
     "evidence": "Browne signed for the Company"},
    {"mention": "%(prefix)s Howqua", "normalized_name": "%(prefix)s Howqua",
     "role": "security merchant", "mention_type": "person_name", "confidence": 0.9,
     "evidence": "Howqua agreed to the contract"}
  ],
  "relationships": [
    {"source_mention": "%(prefix)s Henry Browne", "target_mention": "%(prefix)s Howqua",
     "relation_type": "commercial", "relation_label": "woollen contract",
     "directionality": "directed", "confidence": 0.8,
     "evidence": "Browne and Howqua settled the woollen contract"}
  ]
}
"""


def _metrics():
    from app.schemas.kg_extraction_schema import ExtractionTaskMetrics

    return ExtractionTaskMetrics(
        model_name="auto-mock",
        latency_ms=1.0,
        prompt_tokens=3,
        completion_tokens=3,
        total_tokens=6,
    )


def _mock_llm():
    """Model is mocked; everything downstream (parser/writer/ledger) stays real."""
    return mock.patch.object(
        kg_llm_client,
        "call_extraction_llm",
        new=AsyncMock(return_value=(EXTRACTION_JSON % {"prefix": SYNTH}, _metrics())),
    )


def _docx_bytes(*paragraphs: str) -> bytes:
    document = DocxDocument()
    document.add_heading(SYNTH + " Consultation", level=1)
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@requires_postgres_test_db
class TestKgAutoExtractionPostgres(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.engine = make_test_engine()
        cls.db = make_session_factory(cls.engine)()
        app.dependency_overrides[get_db] = lambda: cls.db
        cls.client = TestClient(app)
        cls.service = ResearchGroupService(cls.db)
        cls.dispatched = []
        cls._cleanup()
        cls.group_main = cls.service.create_group(GroupCreate(name=GROUP_MAIN))
        cls.group_other = cls.service.create_group(GroupCreate(name=GROUP_OTHER))
        cls.group_history = cls.service.create_group(GroupCreate(name=GROUP_HISTORY))
        cls.counter = 0

    @classmethod
    def tearDownClass(cls):
        # Releasing the session/engine matters: an abandoned open transaction would
        # block the next suite's drop_test_schema().
        try:
            cls._cleanup()
        finally:
            app.dependency_overrides.pop(get_db, None)
            cls.db.close()
            cls.engine.dispose()

    # ------------------------------------------------------------- fixtures
    @classmethod
    def _dispatch_recorder(cls):
        def record(run_id: int) -> None:
            cls.dispatched.append(run_id)

        return record

    @classmethod
    def _import(cls, group_id: int, *files):
        """Import through the real multipart endpoint, with a recording dispatcher."""
        cls.dispatched = []
        payload = [
            (
                "files",
                (
                    name,
                    content,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                ),
            )
            for name, content in files
        ]
        with mock.patch.object(
            ResearchGroupService, "auto_dispatch", cls._dispatch_recorder()
        ):
            return cls.client.post("/api/groups/%d/import" % group_id, files=payload)

    @classmethod
    def _auto_on(cls, budget_blocks: int = 50, budget_tokens: int = 0):
        return mock.patch.multiple(
            settings,
            kg_auto_extract_on_import=True,
            kg_auto_extract_budget_blocks=budget_blocks,
            kg_auto_extract_budget_tokens=budget_tokens,
        )

    @classmethod
    def _auto_off(cls):
        return mock.patch.object(settings, "kg_auto_extract_on_import", False)

    @classmethod
    def _runs(cls, group_id: int):
        return (
            cls.db.query(KGExtractionRun)
            .filter(KGExtractionRun.group_id == group_id)
            .order_by(KGExtractionRun.id.asc())
            .all()
        )

    @classmethod
    def _process(cls, run_id: int):
        """Simulate the worker: the default extractor (real pipeline, mocked model)."""
        with _mock_llm():
            return extraction_run_service.run_until_stopped(
                cls.db, run_id, batch_size=5, max_batches=20
            )

    @classmethod
    def _cleanup(cls):
        """Remove only what these tests created, identified by their own blocks/runs.

        Persons and relationships are discovered through the mentions/evidence of the
        synthetic documents, so cleanup does not depend on how the writer normalized
        a name.
        """
        db = cls.db
        try:
            doc_ids = [
                row[0]
                for row in db.query(Document.id)
                .filter(Document.original_filename.like("Synthetic_AUTO_%"))
                .all()
            ]
            page_ids = [
                row[0]
                for row in db.query(Page.id).filter(Page.document_id.in_(doc_ids or [-1])).all()
            ]
            block_ids = [
                row[0]
                for row in db.query(Block.id).filter(Block.page_id.in_(page_ids or [-1])).all()
            ]
            person_ids = {
                row[0]
                for row in db.query(PersonMention.person_id)
                .filter(PersonMention.block_id.in_(block_ids or [-1]))
                .all()
                if row[0] is not None
            }
            relationship_ids = {
                row[0]
                for row in db.query(RelationshipEvidence.relationship_id)
                .filter(RelationshipEvidence.block_id.in_(block_ids or [-1]))
                .all()
            }
            if relationship_ids:
                db.query(RelationshipEvidence).filter(
                    RelationshipEvidence.relationship_id.in_(relationship_ids)
                ).delete(synchronize_session=False)
                db.query(Relationship).filter(
                    Relationship.id.in_(relationship_ids)
                ).delete(synchronize_session=False)
            if block_ids:
                db.query(PersonMention).filter(
                    PersonMention.block_id.in_(block_ids)
                ).delete(synchronize_session=False)
            for person_id in person_ids:
                remaining = (
                    db.query(PersonMention.id)
                    .filter(PersonMention.person_id == person_id)
                    .first()
                )
                linked = (
                    db.query(Relationship.id)
                    .filter(
                        (Relationship.source_person_id == person_id)
                        | (Relationship.target_person_id == person_id)
                    )
                    .first()
                )
                if remaining is None and linked is None:
                    db.query(Person).filter(Person.id == person_id).delete(
                        synchronize_session=False
                    )
            if block_ids:
                db.query(Block).filter(Block.id.in_(block_ids)).delete(
                    synchronize_session=False
                )
            if page_ids:
                db.query(Page).filter(Page.id.in_(page_ids)).delete(
                    synchronize_session=False
                )
            if doc_ids:
                db.query(Document).filter(Document.id.in_(doc_ids)).delete(
                    synchronize_session=False
                )

            group_ids = [
                row[0]
                for row in db.query(ResearchGroup.id)
                .filter(ResearchGroup.name.like(SYNTH + "%"))
                .all()
            ]
            if group_ids:
                run_ids = [
                    row[0]
                    for row in db.query(KGExtractionRun.id)
                    .filter(KGExtractionRun.group_id.in_(group_ids))
                    .all()
                ]
                if run_ids:
                    db.query(KGExtractionItem).filter(
                        KGExtractionItem.run_id.in_(run_ids)
                    ).delete(synchronize_session=False)
                    db.query(KGExtractionRun).filter(
                        KGExtractionRun.id.in_(run_ids)
                    ).delete(synchronize_session=False)
                db.query(ResearchGroupDocument).filter(
                    ResearchGroupDocument.group_id.in_(group_ids)
                ).delete(synchronize_session=False)
                db.query(ResearchGroup).filter(ResearchGroup.id.in_(group_ids)).delete(
                    synchronize_session=False
                )
            db.commit()
        except Exception:  # noqa: BLE001 - cleanup must not mask failures
            db.rollback()

    def _name(self, stem: str) -> str:
        type(self).counter += 1
        return "Synthetic_AUTO_%s_%d.docx" % (stem, self.counter)

    def _fresh_group(self, name: str) -> ResearchGroup:
        return self.service.create_group(GroupCreate(name=name))

    # ------------------------------------------------------------ behaviour
    def test_01_creating_a_group_starts_nothing(self):
        with self._auto_on():
            group = self._fresh_group(SYNTH + " Empty group")
            detail = self.client.get("/api/groups/%d?refresh=false" % group.id).json()
            self.assertEqual(detail["auto_research"]["effective"], True)
            self.assertEqual(detail["auto_research"]["reason"], "started")
            self.assertEqual(self._runs(group.id), [])
            scope = self.client.get("/api/kg/groups/%d/graph/scope" % group.id).json()
            self.assertEqual(scope["extraction_state"], "no_documents")
            self.assertIsNone(scope["extraction_run"])
            self.assertTrue(scope["auto_research"]["effective"])
            self.assertEqual(self.dispatched, [])

    def test_02_single_word_import_auto_queues_and_builds_the_group_graph(self):
        with self._auto_on(budget_blocks=50):
            group = self._fresh_group(SYNTH + " Single import")
            body = (
                SYNTH
                + " 1792: Browne and Howqua settled the woollen contract in Canton."
            )
            response = self._import(group.id, (self._name("single"), _docx_bytes(body)))
            self.assertEqual(response.status_code, 200, response.text)
            payload = response.json()
            self.assertEqual(payload["accepted_count"], 1)
            auto = payload["auto_research"][0]
            self.assertTrue(auto["started"], auto)
            self.assertEqual(auto["reason"], "started")
            self.assertEqual(auto["dispatch_state"], "dispatched")
            run_id = auto["run_id"]
            self.assertIn(run_id, self.dispatched)

            run = self.db.query(KGExtractionRun).filter(KGExtractionRun.id == run_id).one()
            self.assertEqual(run.trigger, KGExtractionRunTrigger.AUTO_IMPORT.value)
            self.assertEqual(run.dispatch_state, "dispatched")
            self.assertGreaterEqual(run.total_blocks, 2)
            document_id = payload["items"][0]["document_id"]
            self.assertEqual(run.document_id, document_id)

            # Status is visible before the worker runs.
            queued = self.client.get("/api/kg/groups/%d/graph/scope" % group.id).json()
            self.assertEqual(queued["extraction_state"], "queued_extraction")
            self.assertEqual(queued["extraction_run"]["trigger"], "auto_import")

            snapshot = self._process(run_id)
            self.assertIn(snapshot["status"], ("completed", "partial"), snapshot)
            self.assertGreater(snapshot["processed_blocks"], 0)

            scope = self.client.get("/api/kg/groups/%d/graph/scope" % group.id).json()
            self.assertEqual(scope["extraction_state"], "ready")
            self.assertEqual(scope["extraction_run"]["status"], "completed")
            self.assertGreaterEqual(scope["mention_count"], 2)
            self.assertEqual(scope["relationship_count_official"], 0)
            self.assertGreaterEqual(scope["relationship_count_candidates"], 1)

            overview = self.client.get(
                "/api/kg/groups/%d/graph/overview?view=all" % group.id
            ).json()
            self.assertEqual(len(overview["persons"]), 2)
            self.assertEqual(len(overview["relationships"]), 1)
            relationship_id = overview["relationships"][0]["id"]

            # Evidence of that relationship is an in-group block of the imported Word.
            evidence = self.client.get(
                "/api/kg/groups/%d/relationships/%d/evidence"
                % (group.id, relationship_id)
            ).json()
            self.assertGreaterEqual(len(evidence["items"]), 1)
            self.assertEqual(evidence["items"][0]["document_id"], document_id)

            # Extraction never accepts anything by itself.
            official = self.client.get(
                "/api/kg/groups/%d/graph/overview?view=official" % group.id
            ).json()
            self.assertEqual(official["persons"], [])
            self.assertEqual(official["relationships"], [])
            for item in overview["persons"]:
                self.assertEqual(item["identity_status"], IdentityStatus.PROVISIONAL.value)
            relationship = (
                self.db.query(Relationship)
                .filter(Relationship.id == relationship_id)
                .one()
            )
            self.assertEqual(relationship.claim_status, ClaimStatus.ASSERTED)
            self.assertEqual(relationship.review_status, ReviewStatusKG.UNREVIEWED)

    def test_03_batch_import_queues_every_document(self):
        with self._auto_on():
            group = self._fresh_group(SYNTH + " Batch import")
            files = [
                (self._name("batch1"), _docx_bytes(SYNTH + " first batch document")),
                (self._name("batch2"), _docx_bytes(SYNTH + " second batch document")),
                (self._name("batch3"), _docx_bytes(SYNTH + " third batch document")),
            ]
            response = self._import(group.id, *files)
            self.assertEqual(response.status_code, 200, response.text)
            payload = response.json()
            self.assertEqual(payload["accepted_count"], 3)
            self.assertEqual(payload["failed_count"], 0)
            started = [item for item in payload["auto_research"] if item["started"]]
            self.assertEqual(len(started), 3, payload["auto_research"])
            run_ids = {item["run_id"] for item in started}
            self.assertEqual(len(run_ids), 3)
            self.assertEqual(sorted(self.dispatched), sorted(run_ids))
            runs = self._runs(group.id)
            self.assertEqual(len(runs), 3)
            self.assertEqual({run.trigger for run in runs}, {"auto_import"})
            # One run per document: no document is silently dropped or shared.
            self.assertEqual(len({run.document_id for run in runs}), 3)

    def test_04_duplicate_import_does_not_duplicate_extraction(self):
        with self._auto_on():
            group = self._fresh_group(SYNTH + " Duplicate import")
            name = self._name("dup")
            content = _docx_bytes(SYNTH + " duplicate content document")
            first = self._import(group.id, (name, content)).json()
            self.assertEqual(first["accepted_count"], 1)
            run_id = first["auto_research"][0]["run_id"]
            self._process(run_id)
            mentions_before = self.db.query(PersonMention.id).count()

            second = self._import(group.id, (name, content)).json()
            self.assertEqual(second["duplicate_count"], 1)
            self.assertEqual(second["accepted_count"], 0)
            self.assertEqual(second["auto_research"], [])
            self.assertEqual(len(self._runs(group.id)), 1)
            self.assertEqual(self.db.query(PersonMention.id).count(), mentions_before)

    def test_05_queue_unavailable_keeps_the_run_and_can_be_redelivered(self):
        with self._auto_on():
            group = self._fresh_group(SYNTH + " Queue down")
            name = self._name("queuedown")
            content = _docx_bytes(SYNTH + " queue unavailable document")

            def broken(run_id: int) -> None:
                raise ConnectionError("broker unreachable")

            with mock.patch.object(ResearchGroupService, "auto_dispatch", broken):
                response = self.client.post(
                    "/api/groups/%d/import" % group.id,
                    files=[("files", (name, content, "application/octet-stream"))],
                )
            # The import itself succeeded: a broken queue never fails the upload.
            self.assertEqual(response.status_code, 200, response.text)
            payload = response.json()
            self.assertEqual(payload["accepted_count"], 1)
            auto = payload["auto_research"][0]
            self.assertTrue(auto["started"])
            self.assertEqual(auto["dispatch_state"], "failed")
            self.assertIn("broker unreachable", auto["dispatch_error"])

            run_id = auto["run_id"]
            run = self.db.query(KGExtractionRun).filter(KGExtractionRun.id == run_id).one()
            self.assertEqual(run.dispatch_state, "failed")
            self.assertEqual(run.dispatch_attempts, 1)

            scope = self.client.get("/api/kg/groups/%d/graph/scope" % group.id).json()
            self.assertEqual(scope["extraction_run"]["dispatch_state"], "failed")
            self.assertIn("broker unreachable", scope["extraction_run"]["dispatch_error"])

            undelivered = self.client.get("/api/kg/extraction/dispatch-queue").json()
            self.assertIn(run_id, [item["run_id"] for item in undelivered])

            with mock.patch.object(
                ResearchGroupService, "auto_dispatch", self._dispatch_recorder()
            ), mock.patch.object(
                auto_extraction, "_default_dispatch", self._dispatch_recorder()
            ):
                retried = self.client.post(
                    "/api/kg/extraction/runs/%d/dispatch" % run_id
                )
            self.assertEqual(retried.status_code, 200, retried.text)
            self.assertEqual(retried.json()["dispatch_state"], "dispatched")
            self.assertIsNone(retried.json()["dispatch_error"])
            self.assertNotIn(
                run_id,
                [
                    item["run_id"]
                    for item in self.client.get(
                        "/api/kg/extraction/dispatch-queue"
                    ).json()
                ],
            )

    def test_06_new_document_does_not_reextract_the_old_blocks(self):
        with self._auto_on():
            group = self._fresh_group(SYNTH + " Incremental")
            first = self._import(
                group.id, (self._name("inc1"), _docx_bytes(SYNTH + " first document body"))
            ).json()
            first_run = first["auto_research"][0]["run_id"]
            first_doc = first["items"][0]["document_id"]
            self._process(first_run)
            first_block_ids = {
                row[0]
                for row in self.db.query(Block.id)
                .join(Page, Page.id == Block.page_id)
                .filter(Page.document_id == first_doc)
                .all()
            }
            first_mentions = {
                row[0]
                for row in self.db.query(PersonMention.block_id)
                .filter(PersonMention.block_id.in_(first_block_ids))
                .all()
            }

            second = self._import(
                group.id, (self._name("inc2"), _docx_bytes(SYNTH + " second document body"))
            ).json()
            second_run_id = second["auto_research"][0]["run_id"]
            second_run = (
                self.db.query(KGExtractionRun)
                .filter(KGExtractionRun.id == second_run_id)
                .one()
            )
            self.assertNotEqual(second_run_id, first_run)
            second_doc = second["items"][0]["document_id"]
            second_block_count = (
                self.db.query(Block.id)
                .join(Page, Page.id == Block.page_id)
                .filter(Page.document_id == second_doc)
                .count()
            )
            # The new run covers only the new document, never the whole group again.
            self.assertEqual(second_run.total_blocks, second_block_count)
            self._process(second_run_id)

            # Blocks of the first document were not touched again.
            first_items = (
                self.db.query(KGExtractionItem)
                .filter(
                    KGExtractionItem.run_id == second_run_id,
                    KGExtractionItem.block_id.in_(first_block_ids),
                )
                .all()
            )
            self.assertEqual(first_items, [])
            self.assertEqual(
                {
                    row[0]
                    for row in self.db.query(PersonMention.block_id)
                    .filter(PersonMention.block_id.in_(first_block_ids))
                    .all()
                },
                first_mentions,
            )

    def test_07_budget_pause_and_resume_on_an_automatic_run(self):
        with self._auto_on(budget_blocks=1):
            group = self._fresh_group(SYNTH + " Budget")
            body = _docx_bytes(
                SYNTH + " paragraph one", SYNTH + " paragraph two", SYNTH + " paragraph three"
            )
            payload = self._import(group.id, (self._name("budget"), body)).json()
            run_id = payload["auto_research"][0]["run_id"]
            run = self.db.query(KGExtractionRun).filter(KGExtractionRun.id == run_id).one()
            self.assertEqual(run.budget_blocks, 1)
            snapshot = self._process(run_id)
            self.assertEqual(snapshot["status"], "paused_budget", snapshot)
            self.assertIn("预算", snapshot["pause_reason"])

            scope = self.client.get("/api/kg/groups/%d/graph/scope" % group.id).json()
            self.assertEqual(scope["extraction_state"], "paused_extraction")
            self.assertIn("预算", scope["state_message"])

            refused = self.client.post(
                "/api/kg/extraction/runs/%d/resume?async_task=false" % run_id
            )
            self.assertEqual(refused.status_code, 409, refused.text)

            resumed = self.client.post(
                "/api/kg/extraction/runs/%d/resume?async_task=false&budget_blocks=50"
                % run_id
            )
            self.assertEqual(resumed.status_code, 200, resumed.text)
            final = self._process(run_id)
            self.assertIn(final["status"], ("completed", "partial"), final)
            self.assertEqual(final["attempted_blocks"], final["total_blocks"])

    def test_08_two_groups_never_share_extracted_evidence(self):
        with self._auto_on():
            left = self._fresh_group(SYNTH + " Isolation left")
            right = self._fresh_group(SYNTH + " Isolation right")
            left_payload = self._import(
                left.id, (self._name("left"), _docx_bytes(SYNTH + " left group text"))
            ).json()
            right_payload = self._import(
                right.id, (self._name("right"), _docx_bytes(SYNTH + " right group text"))
            ).json()
            self._process(left_payload["auto_research"][0]["run_id"])

            left_overview = self.client.get(
                "/api/kg/groups/%d/graph/overview?view=all" % left.id
            ).json()
            right_overview = self.client.get(
                "/api/kg/groups/%d/graph/overview?view=all" % right.id
            ).json()
            self.assertGreaterEqual(len(left_overview["relationships"]), 1)
            # Same global identities exist, but the right group has no evidence yet.
            self.assertEqual(right_overview["relationships"], [])
            self.assertEqual(right_overview["persons"], [])

            self._process(right_payload["auto_research"][0]["run_id"])
            right_after = self.client.get(
                "/api/kg/groups/%d/graph/overview?view=all" % right.id
            ).json()
            self.assertTrue(right_after["relationships"])
            # The same claim may be supported from both groups, but each group must
            # only ever show evidence from its own documents.
            left_document_id = left_payload["items"][0]["document_id"]
            right_document_id = right_payload["items"][0]["document_id"]
            for rel in right_after["relationships"]:
                evidence = self.client.get(
                    "/api/kg/groups/%d/relationships/%d/evidence" % (right.id, rel["id"])
                ).json()
                self.assertTrue(evidence["items"])
                self.assertEqual(
                    {item["document_id"] for item in evidence["items"]},
                    {right_document_id},
                )
            for rel in left_overview["relationships"]:
                evidence = self.client.get(
                    "/api/kg/groups/%d/relationships/%d/evidence" % (left.id, rel["id"])
                ).json()
                self.assertTrue(evidence["items"])
                self.assertEqual(
                    {item["document_id"] for item in evidence["items"]},
                    {left_document_id},
                )

    def test_09_enabling_the_switch_never_backfills_history(self):
        name = self._name("history")
        content = _docx_bytes(SYNTH + " imported while auto research was off")
        with self._auto_off():
            group = self._fresh_group(SYNTH + " History only group")
            payload = self._import(group.id, (name, content)).json()
            self.assertEqual(payload["accepted_count"], 1)
            self.assertEqual(payload["auto_research"][0]["reason"], "global_disabled")
            self.assertFalse(payload["auto_research"][0]["started"])
            self.assertEqual(self._runs(group.id), [])
            document_id = payload["items"][0]["document_id"]

        with self._auto_on():
            # Turning the switch on (and reading pages) must not create work.
            detail = self.client.get("/api/groups/%d?refresh=true" % group.id).json()
            self.assertTrue(detail["auto_research"]["effective"])
            self.client.get("/api/kg/groups/%d/graph/scope" % group.id)
            self.client.get("/api/kg/groups/%d/extraction" % group.id)
            self.service.update_group(group.id, _auto_toggle(True))
            self.assertEqual(self._runs(group.id), [])

            # Only a real event queues: a newly attached document.
            with mock.patch.object(
                ResearchGroupService, "auto_dispatch", self._dispatch_recorder()
            ):
                attached = self.client.post(
                    "/api/groups/%d/documents" % group.id,
                    json={"document_ids": [document_id]},
                )
            # It is already a member: duplicate, therefore no research event.
            self.assertEqual(attached.json()["added"], [])
            self.assertEqual(self._runs(group.id), [])

            fresh = self._import(
                group.id, (self._name("history2"), _docx_bytes(SYNTH + " new material"))
            ).json()
            self.assertEqual(len(fresh["auto_research"]), 1)
            self.assertTrue(fresh["auto_research"][0]["started"])
            self.assertEqual(len(self._runs(group.id)), 1)

    def test_10_attaching_an_existing_word_is_a_research_event(self):
        name = self._name("attach")
        content = _docx_bytes(SYNTH + " uploaded first, attached later")
        upload = self.client.post(
            "/api/documents/upload",
            files=[("file", (name, content, "application/octet-stream"))],
        )
        self.assertEqual(upload.status_code, 200, upload.text)
        document_id = upload.json()["id"]
        with self._auto_off():
            group = self._fresh_group(SYNTH + " Attach group")
            self.assertEqual(self._runs(group.id), [])
        with self._auto_on():
            self.dispatched = []
            type(self).dispatched = []
            with mock.patch.object(
                ResearchGroupService, "auto_dispatch", self._dispatch_recorder()
            ):
                response = self.client.post(
                    "/api/groups/%d/documents" % group.id,
                    json={"document_ids": [document_id]},
                )
            self.assertEqual(response.status_code, 200, response.text)
            body = response.json()
            self.assertEqual(body["added"], [document_id])
            self.assertEqual(len(body["auto_research"]), 1)
            outcome = body["auto_research"][0]
            self.assertTrue(outcome["started"], outcome)
            self.assertEqual(outcome["document_id"], document_id)
            self.assertIn(outcome["run_id"], type(self).dispatched)
            run = (
                self.db.query(KGExtractionRun)
                .filter(KGExtractionRun.id == outcome["run_id"])
                .one()
            )
            self.assertEqual(run.trigger, KGExtractionRunTrigger.AUTO_ATTACH.value)

    def test_11_group_switch_and_deployment_switch_are_reported_separately(self):
        with self._auto_on():
            group = self._fresh_group(SYNTH + " Toggle group")
            self.service.update_group(group.id, _auto_toggle(False))
            detail = self.client.get("/api/groups/%d?refresh=false" % group.id).json()
            self.assertFalse(detail["group"]["auto_extract"])
            self.assertEqual(detail["auto_research"]["reason"], "group_disabled")
            self.assertFalse(detail["auto_research"]["effective"])
            payload = self._import(
                group.id, (self._name("toggle"), _docx_bytes(SYNTH + " disabled group"))
            ).json()
            # The reason is reported per document instead of silently queueing.
            self.assertEqual(len(payload["auto_research"]), 1)
            self.assertEqual(payload["auto_research"][0]["reason"], "group_disabled")
            self.assertFalse(payload["auto_research"][0]["started"])
            self.assertEqual(self._runs(group.id), [])

        with self._auto_off():
            detail = self.client.get("/api/groups/%d?refresh=false" % group.id).json()
            self.assertEqual(detail["auto_research"]["reason"], "global_disabled")

    def test_13_import_while_a_group_run_is_active_is_not_dropped(self):
        """Scenario D: a running group-level run must not swallow a new document."""
        with self._auto_on():
            group = self._fresh_group(SYNTH + " Concurrent group run")
            first = self._import(
                group.id, (self._name("conc1"), _docx_bytes(SYNTH + " first concurrent doc"))
            ).json()
            self._process(first["auto_research"][0]["run_id"])

            # A manual, group-scoped run is now queued and stays active.
            manual = self.client.post(
                "/api/kg/groups/%d/extraction?async_task=false&budget_blocks=50" % group.id
            )
            self.assertEqual(manual.status_code, 201, manual.text)
            manual_id = manual.json()["run_id"]

            # Importing another Word must still be queued (document scope, no 409 loss).
            second = self._import(
                group.id, (self._name("conc2"), _docx_bytes(SYNTH + " second concurrent doc"))
            ).json()
            self.assertEqual(second["accepted_count"], 1)
            auto = second["auto_research"][0]
            self.assertTrue(auto["started"], auto)
            self.assertNotEqual(auto["run_id"], manual_id)

            manual_run = (
                self.db.query(KGExtractionRun)
                .filter(KGExtractionRun.id == manual_id)
                .one()
            )
            # The pre-existing run keeps its own scope and status: nothing was stolen.
            self.assertEqual(manual_run.scope_type, "group")
            self.assertEqual(manual_run.status, "queued")
            self.assertIsNone(manual_run.document_id)

    def test_12_missing_model_configuration_blocks_the_queue(self):
        with self._auto_on(), mock.patch.multiple(
            settings, gpt_api_key="", gpt_base_url=""
        ):
            group = self._fresh_group(SYNTH + " No model")
            payload = self._import(
                group.id, (self._name("nomodel"), _docx_bytes(SYNTH + " no model text"))
            ).json()
            self.assertEqual(payload["auto_research"][0]["reason"], "model_not_configured")
            self.assertFalse(payload["auto_research"][0]["started"])
            self.assertEqual(self._runs(group.id), [])


def _auto_toggle(value: bool):
    from app.schemas.research_group import GroupUpdate

    return GroupUpdate(auto_extract=value)
