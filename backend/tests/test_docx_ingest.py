"""Unit tests for Word (.docx) text extraction (no database, no file system).

The extraction is the part that decides what becomes searchable archival text, so
it is tested as a pure function over bytes: paragraph order, heading detection,
table rows, and the explicit failure modes (not-a-docx, empty text).
"""
import unittest
from io import BytesIO

from docx import Document as DocxDocument

from app.services.docx_service import (
    DocxParseError,
    extract_paragraphs,
    to_markdown,
)


def _docx_bytes(build) -> bytes:
    document = DocxDocument()
    build(document)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class TestDocxExtraction(unittest.TestCase):
    def test_paragraphs_keep_reading_order_and_skip_blanks(self):
        def build(document):
            document.add_heading("Canton Consultations 1792", level=1)
            document.add_paragraph("On 21 November 1792 the woollen contract was settled.")
            document.add_paragraph("")
            document.add_paragraph("    ")
            document.add_paragraph("Signed by Browne and Howqua.")

        paragraphs = extract_paragraphs(_docx_bytes(build))

        self.assertEqual(
            [p.text for p in paragraphs],
            [
                "Canton Consultations 1792",
                "On 21 November 1792 the woollen contract was settled.",
                "Signed by Browne and Howqua.",
            ],
        )
        self.assertEqual([p.index for p in paragraphs], [0, 1, 2])
        self.assertTrue(paragraphs[0].is_heading)
        self.assertFalse(paragraphs[1].is_heading)

    def test_chinese_paragraphs_are_preserved_verbatim(self):
        text = "乾隆五十七年十一月，大班 Browne 与行商 Howqua 议定毛织品合同。"

        def build(document):
            document.add_paragraph(text)

        paragraphs = extract_paragraphs(_docx_bytes(build))
        self.assertEqual(paragraphs[0].text, text)

    def test_table_rows_become_text_blocks(self):
        def build(document):
            document.add_paragraph("Goods delivered")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Item"
            table.cell(0, 1).text = "Quantity"
            table.cell(1, 0).text = "Woollens"
            table.cell(1, 1).text = "1240 pieces"

        paragraphs = extract_paragraphs(_docx_bytes(build))
        texts = [p.text for p in paragraphs]
        self.assertIn("Woollens | 1240 pieces", texts)
        self.assertEqual(paragraphs[-1].style, "Table")
        self.assertEqual(paragraphs[-1].index, len(paragraphs) - 1)

    def test_not_a_docx_is_reported_not_silently_empty(self):
        with self.assertRaises(DocxParseError):
            extract_paragraphs(b"%PDF-1.4 not a word package")
        with self.assertRaises(DocxParseError):
            extract_paragraphs(b"")

    def test_document_without_text_is_rejected(self):
        def build(document):
            document.add_paragraph("")
            document.add_paragraph("   ")

        with self.assertRaises(DocxParseError) as ctx:
            extract_paragraphs(_docx_bytes(build))
        self.assertIn("没有可导入的文本", str(ctx.exception))

    def test_markdown_keeps_headings_and_paragraph_breaks(self):
        def build(document):
            document.add_heading("Section A", level=1)
            document.add_paragraph("First line.")
            document.add_paragraph("第二段。")

        markdown = to_markdown(extract_paragraphs(_docx_bytes(build)))
        self.assertTrue(markdown.startswith("# Section A"))
        self.assertIn("First line.", markdown)
        self.assertIn("第二段。", markdown)
        self.assertTrue(markdown.endswith("\n"))


if __name__ == "__main__":
    unittest.main()
